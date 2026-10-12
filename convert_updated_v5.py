import os
import re
import subprocess
import sys
from datetime import datetime

import pypandoc
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.run import Run

# Needs: pip install pypandoc python-docx
# with time tracker + Times New Roman + black text + A4 page size + page numbers

FONT_NAME = "Times New Roman"
BLACK = RGBColor(0, 0, 0)

THEME_FONT_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")


def strip_boxed(text: str) -> str:
    """Remove \boxed{...} wrapper, keeping inner content (handles nesting)."""
    result = []
    i = 0
    while i < len(text):
        if text[i:i+7] == r'\boxed{':
            depth = 1
            j = i + 7
            start = j
            while j < len(text) and depth > 0:
                if text[j] == '{':
                    depth += 1
                elif text[j] == '}':
                    depth -= 1
                j += 1
            inner = text[start:j-1]
            result.append(strip_boxed(inner))  # handle nested \boxed too
            i = j
        else:
            result.append(text[i])
            i += 1
    return ''.join(result)


def fix_rm_command(text: str) -> str:
    r"""Convert old-style {\rm TEXT} to modern \mathrm{TEXT}, which Pandoc supports."""
    text = re.sub(r'\{\\rm\s+([^{}]*)\}', r'\\mathrm{\1}', text)
    text = text.replace(r'\rm ', r'\mathrm ')
    return text


def fix_bracket_equations(text: str) -> tuple[str, int]:
    r"""
    ChatGPT sometimes wraps equations in bare [ ... ] on their own lines
    instead of \[ ... \]. Convert real content into proper \[ \] math
    delimiters so Pandoc parses it. Flag genuinely empty/lost ones instead
    of silently dropping them.
    r"""
    broken_count = 0

    def replacer(match):
        nonlocal broken_count
        inner = match.group(1).strip()
        # strip stray bullet/dot characters some copies leave at line starts
        inner = re.sub(r'^[·•\-\*]\s*', '', inner, flags=re.MULTILINE).strip()

        letters = re.sub(r'[^a-zA-Z0-9]', '', inner)
        if len(letters) <= 3:
            broken_count += 1
            return f'**[⚠ BROKEN EQUATION — original content lost: "{inner}"]**'

        # real content -> wrap as proper LaTeX display math
        return f"\\[{inner}\\]"

    text = re.sub(r'^\[\s*\n(.*?)\n\]\s*$', replacer, text, flags=re.MULTILINE | re.DOTALL)
    return text, broken_count


def clean_text(text: str) -> tuple[str, int]:
    # Remove ChatGPT's image_group metadata lines
    text = re.sub(r'^image_group\{.*?\}$', '', text, flags=re.MULTILINE)
    # Strip \boxed{...} wrappers
    text = strip_boxed(text)
    # Fix old-style {\rm TEXT} font commands
    text = fix_rm_command(text)
    # Rescue bare [ ... ] equation blocks into proper \[ \] math
    text, broken_count = fix_bracket_equations(text)
    return text, broken_count


def generate_output_filename(base_name: str = "Converted_v5", ext: str = "docx") -> str:
    """Generate a filename stamped with the current date and time."""
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    return f"{base_name}_{timestamp}.{ext}"


# ----------------------------------------------------------------------------
# Post-processing: Times New Roman, black text, A4 page
# ----------------------------------------------------------------------------

def _force_rfonts(rfonts):
    """Point every font slot at Times New Roman and drop theme-font overrides
    (theme attributes would otherwise win over the explicit font name)."""
    for attr in THEME_FONT_ATTRS:
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{slot}"), FONT_NAME)


def _style_all_styles(doc):
    """Set font + black color on every style; make headings/titles bold."""
    # document-wide defaults
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr_default = defaults.find(qn("w:rPrDefault"))
        if rpr_default is not None:
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is not None:
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = OxmlElement("w:rFonts")
                    rpr.insert(0, rfonts)
                _force_rfonts(rfonts)

    for style in doc.styles:
        try:
            style.font.name = FONT_NAME
            style.font.color.rgb = BLACK
            rfonts = style.element.rPr.rFonts
            if rfonts is not None:
                _force_rfonts(rfonts)

            name = (style.name or "").lower()
            if name.startswith("heading") or name in ("title", "subtitle"):
                style.font.bold = True
        except Exception:
            # some style types (e.g. numbering) don't support font formatting
            continue


def _style_all_runs(doc):
    """Direct run formatting (hyperlinks, inline overrides) must be fixed too."""
    parts = [doc.element.body]
    for section in doc.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            try:
                parts.append(hf._element)
            except Exception:
                pass

    for root in parts:
        for r in root.iter(qn("w:r")):
            run = Run(r, None)
            run.font.name = FONT_NAME
            run.font.color.rgb = BLACK
            rfonts = r.rPr.rFonts
            if rfonts is not None:
                _force_rfonts(rfonts)


def _set_a4(doc):
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(25.4)
        section.right_margin = Mm(25.4)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)


PAGE_NUMBER_SIZE_PT = 10


def _field_run(paragraph, kind=None, instr=None, text=None):
    """Add one run (Times New Roman, 10 pt, black) holding a Word field part."""
    run = paragraph.add_run()
    run.font.name = FONT_NAME
    run.font.size = Pt(PAGE_NUMBER_SIZE_PT)
    run.font.color.rgb = BLACK
    _force_rfonts(run._r.get_or_add_rPr().get_or_add_rFonts())
    if kind:
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), kind)
        run._r.append(fld)
    if instr:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = instr
        run._r.append(it)
    if text:
        run.text = text
    return run


def _add_page_numbers(doc):
    """Centered page number at the bottom of every page (Times New Roman, 10 pt)."""
    for section in doc.sections:
        section.footer.is_linked_to_previous = False
        footer = section.footer
        # clear anything already in the footer, keep one empty paragraph
        for p in footer.paragraphs[1:]:
            p._element.getparent().remove(p._element)
        para = footer.paragraphs[0]
        for r in list(para.runs):
            r._r.getparent().remove(r._r)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        _field_run(para, kind="begin")
        _field_run(para, instr=" PAGE ")
        _field_run(para, kind="separate")
        _field_run(para, text="1")
        _field_run(para, kind="end")


def format_docx(path: str, page_numbers: bool = True):
    doc = Document(path)
    _style_all_styles(doc)
    _style_all_runs(doc)
    _set_a4(doc)
    if page_numbers:
        _add_page_numbers(doc)
    doc.save(path)


# ----------------------------------------------------------------------------

def open_file(path: str):
    """Open the finished file with its default program (Word for .docx)."""
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)                      # Windows -> opens in Word
        elif sys.platform == "darwin":
            subprocess.run(["open", path], check=False)   # macOS
        else:
            subprocess.run(["xdg-open", path], check=False)  # Linux
    except Exception as e:
        print(f"Could not open the file automatically: {e}")


def markdown_to_docx(input_text: str, output_path: str = None, auto_open: bool = True):
    if output_path is None:
        output_path = generate_output_filename()
    output_path = os.path.abspath(output_path)

    cleaned, broken_count = clean_text(input_text)
    pypandoc.convert_text(
        cleaned,
        to="docx",
        format="markdown+tex_math_single_backslash+tex_math_dollars",
        outputfile=output_path,
        extra_args=["--standalone"]
    )

    # Apply Times New Roman, black color, bold headings and A4 size
    format_docx(output_path)

    print(f"Saved: {output_path}")
    if broken_count:
        print(f"⚠ Warning: {broken_count} equation(s) had missing content in the "
              f"source text and were flagged in the document instead of silently "
              f"dropped. Search the docx for 'BROKEN EQUATION' and fix those by hand.")

    if auto_open:
        open_file(output_path)


if __name__ == "__main__":
    print("\033[93m===================================\033[0m")
    print("\033[91mThis Program is Created by:\nMd.Sabbir Hossain\nDepartment of Physics,PUST.\033[0m")
    print("\033[93m===================================\033[0m")
    print("\033[94mPaste your AI text below. Press Enter,then\033[0m \033[95mCtrl+Z\033[0m \033[94mthen Enter (Windows) when done:\033[0m")
    text = sys.stdin.read()
    markdown_to_docx(text)
